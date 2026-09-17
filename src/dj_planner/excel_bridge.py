from __future__ import annotations

from datetime import datetime, time
from typing import List
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .planner_core import canonical_dj, HistorySet, ScheduleRow


def _clock(v):
    if v is None:
        return None
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    if isinstance(v, (float, int)) and 0 <= float(v) < 1:
        mins = int(round(float(v) * 1440)) % 1440
        return f"{mins // 60:02d}:{mins % 60:02d}"
    s = str(v).strip()
    return s[-8:] if ":" in s and " " in s else s


def import_v2(path: str):
    wb = load_workbook(path, data_only=False)

    roster = []
    if "DJs" in wb.sheetnames:
        for row in wb["DJs"].iter_rows(min_row=2, values_only=True):
            if row[0]:
                roster.append(canonical_dj(row[0]))

    bands = []
    if "Valeurs horaires" in wb.sheetnames:
        for row in wb["Valeurs horaires"].iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[1]:
                continue
            try:
                bands.append((_clock(row[0]), _clock(row[1]), float(row[2] or 0)))
            except Exception:
                pass

    parties = {}
    if "Soirées" in wb.sheetnames:
        for row in wb["Soirées"].iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            parties[str(row[0])] = {
                "id": str(row[0]),
                "name": row[1] or str(row[0]),
                "date": "" if row[2] is None else str(row[2])[:10],
                "start": _clock(row[3]),
                "end": _clock(row[4]),
                "status": row[8] if len(row) > 8 else "",
            }

    history = []
    if "Planning" in wb.sheetnames:
        raw_by_party = {}
        for row in wb["Planning"].iter_rows(min_row=2, values_only=True):
            if not row[0] or not row[4]:
                continue
            pid = str(row[0])
            raw_by_party.setdefault(pid, []).append(row)

        for pid, rows in raw_by_party.items():
            party = parties.get(pid, {"name": pid, "date": ""})
            n = len(rows)
            for i, row in enumerate(rows):
                history.append(
                    HistorySet(
                        party_id=pid,
                        party_name=str(party.get("name", pid)),
                        date=str(party.get("date", "")),
                        dj=canonical_dj(row[4]),
                        start=_clock(row[2]),
                        end=_clock(row[3]),
                        value=float(row[6] or 0),
                        position=0.5 if n <= 1 else i / (n - 1),
                    )
                )

    availability = {}
    if "Disponibilités" in wb.sheetnames:
        ws = wb["Disponibilités"]
        headers = [c.value for c in ws[1]]
        for col_idx, pid in enumerate(headers[1:], start=2):
            if not pid:
                continue
            availability[str(pid)] = {}
            for r in range(2, ws.max_row + 1):
                name = ws.cell(r, 1).value
                if not name:
                    continue
                val = str(ws.cell(r, col_idx).value or "").strip().lower()
                availability[str(pid)][canonical_dj(name)] = val in ("oui", "yes", "1", "true")

    wb.close()
    return {
        "roster": list(dict.fromkeys(roster)),
        "bands": bands,
        "parties": parties,
        "history": history,
        "availability": availability,
    }


def export_schedule(path: str, party: dict, schedule: List[ScheduleRow], history_stats=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    headers = [
        "Event",
        "Date",
        "Start",
        "End",
        "DJ",
        "Duration (min)",
        "Slot value",
        "Locked",
    ]
    ws.append(headers)

    for r in schedule:
        ws.append(
            [
                party.get("name", ""),
                party.get("date", ""),
                r.start,
                r.end,
                r.dj,
                round(r.end_abs - r.start_abs, 2),
                round(r.slot_value, 3),
                "Yes" if r.locked else "No",
            ]
        )

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

    widths = [25, 13, 11, 11, 20, 15, 16, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    if history_stats:
        st = wb.create_sheet("Statistics")
        st.append(["DJ", "Historical sets", "Average slot value", "Average position"])
        for column in "ABCD":
            st.column_dimensions[column].width = 24
        for dj, s in sorted(history_stats.items()):
            st.append(
                [
                    dj,
                    s["sets"],
                    round(s["avg_value"], 3),
                    "" if s["avg_position"] is None else round(s["avg_position"], 3),
                ]
            )
        for cell in st[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")

    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    wb.save(path)
    wb.close()
