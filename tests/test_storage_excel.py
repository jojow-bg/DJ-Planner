from datetime import time

import pytest
from openpyxl import Workbook, load_workbook

from dj_planner.data_store import DataStore
from dj_planner.demo import seed_demo
from dj_planner.excel_bridge import export_schedule, import_v2
from dj_planner.planner_core import DJConstraint, HistorySet, solve_schedule, validate_schedule


def test_new_database_and_persistence(tmp_path):
    path = tmp_path / "nested" / "planner.sqlite3"
    store = DataStore(path)
    assert store.get_djs() == []
    store.add_dj("Nova")
    store.add_dj("Pulse")
    store.append_history([HistorySet("DEMO", "Demo", "2026-01-01", "Nova", "20:00", "21:00")])
    store.remove_dj("Nova")
    store.set_setting("party_start", "21:00")
    store.close()
    store = DataStore(path)
    try:
        assert store.get_djs() == ["Pulse"]
        assert store.get_history(HistorySet)[0].dj == "Nova"
        assert store.get_setting("party_start") == "21:00"
    finally:
        store.close()


def test_demo_is_feasible_and_never_overwrites(tmp_path):
    import json

    path = tmp_path / "demo.sqlite3"
    seed_demo(path)
    with pytest.raises(FileExistsError):
        seed_demo(path)
    store = DataStore(path)
    try:
        states = json.loads(store.get_setting("constraints_json"))
        cs = [DJConstraint(n, c["present"], c["min"], c["max"]) for n, c in states.items()]
        rows = solve_schedule(
            "20:00", "02:00", cs, store.get_bands(), store.get_history(HistorySet)
        )
        assert len(rows) == 6
        assert len(store.get_history(HistorySet)) == 12
        assert validate_schedule("20:00", "02:00", cs, rows) == []
    finally:
        store.close()


def test_excel_exports_actual_duration_and_literal_text(tmp_path):
    rows = solve_schedule(
        "20:00",
        "00:00",
        [DJConstraint(n) for n in ["=Nova", "Pulse", "Echo", "Orbit", "Prism"]],
        [],
    )
    path = tmp_path / "export.xlsx"
    rows[0].locked = True
    stats = {"Nova": {"sets": 2, "avg_value": 3.0, "avg_position": 0.5}}
    export_schedule(path, {"name": "=1+1", "date": "2026-01-01"}, rows, stats)
    wb = load_workbook(path)
    try:
        assert wb.sheetnames == ["Schedule", "Statistics"]
        assert [cell.value for cell in wb.active[1]] == [
            "Event",
            "Date",
            "Start",
            "End",
            "DJ",
            "Duration (min)",
            "Slot value",
            "Locked",
        ]
        assert wb.active["H2"].value == "Yes"
        assert wb.active["H3"].value == "No"
        assert [cell.value for cell in wb["Statistics"][1]] == [
            "DJ",
            "Historical sets",
            "Average slot value",
            "Average position",
        ]
        assert [row[5].value for row in list(wb.active)[1:]] == [45, 45, 45, 45, 60]
        assert wb.active["A2"].data_type == "s"
        assert all(row[4].data_type == "s" for row in list(wb.active)[1:])
    finally:
        wb.close()


def test_legacy_excel_import(tmp_path):
    wb = Workbook()
    wb.active.title = "DJs"
    wb.active.append(["DJ"])
    wb.active.append([" Nova "])
    wb.active.append(["Pulse"])
    ws = wb.create_sheet("Soirées")
    ws.append(["ID", "Name", "Date", "Start", "End"])
    ws.append(["DEMO", "Demo Night", "2026-01-01", time(20), time(22)])
    ws = wb.create_sheet("Disponibilités")
    ws.append(["DJ", "DEMO"])
    ws.append(["Nova", "oui"])
    ws.append(["Pulse", "non"])
    path = tmp_path / "legacy.xlsx"
    wb.save(path)
    wb.close()
    result = import_v2(path)
    assert result["roster"] == ["Nova", "Pulse"]
    assert result["parties"]["DEMO"]["start"] == "20:00"
    assert result["availability"]["DEMO"] == {"Nova": True, "Pulse": False}
